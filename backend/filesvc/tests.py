import datetime
import io
import json
import os
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

from cryptography.hazmat.primitives.asymmetric import rsa
import jwt
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TransactionTestCase
from rest_framework import status
from rest_framework.test import APIClient

from agent.tools import read_file, generate_file, read_generated_spec
from chat.models import Thread, Message, UploadedFile, GeneratedFile
from filesvc.storage import storage
from filesvc.throttles import UploadRateThrottle, ChatRateThrottle
from filesvc.generators.registry import get_generator, summarize_spec

User = get_user_model()

# Key pair for testing JWTs
TEST_PRIVATE_KEY = rsa.generate_private_key(public_exponent=65537, key_size=2048)
TEST_PUBLIC_KEY = TEST_PRIVATE_KEY.public_key()

class MockSigningKey:
    def __init__(self, key):
        self.key = key

def make_test_token(sub="test-user-sub-1", email="user1@example.com"):
    now = datetime.datetime.now(datetime.timezone.utc)
    payload = {
        "sub": sub,
        "email": email,
        "iss": settings.KEYCLOAK_PUBLIC_URL,
        "azp": settings.KEYCLOAK_CLIENT_ID,
        "iat": int(now.timestamp()),
        "exp": int((now + datetime.timedelta(seconds=3600)).timestamp()),
    }
    return jwt.encode(payload, TEST_PRIVATE_KEY, algorithm="RS256", headers={"kid": "test-kid"})

class FileServiceTests(TransactionTestCase):
    """Comprehensive test suite for file uploading, parsing, security, download, and tools."""

    def setUp(self):
        self.client = APIClient()

        # Mock JWKS for Keycloak authentication
        self.mock_jwks = MagicMock()
        self.mock_jwks.get_signing_key_from_jwt.return_value = MockSigningKey(TEST_PUBLIC_KEY)
        self.patcher = patch("chat.authentication.get_jwks_client", return_value=self.mock_jwks)
        self.patcher.start()

        self.token_user1 = make_test_token(sub="sub-user-1", email="user1@example.com")
        self.token_user2 = make_test_token(sub="sub-user-2", email="user2@example.com")

        self.user1 = User.objects.create(username="sub-user-1", email="user1@example.com")
        self.user2 = User.objects.create(username="sub-user-2", email="user2@example.com")

        self.thread1 = Thread.objects.create(user=self.user1, title="Thread 1")
        self.thread2 = Thread.objects.create(user=self.user2, title="Thread 2")

        self.sample_dir = Path(__file__).resolve().parent.parent / "sample_files"

    def tearDown(self):
        self.patcher.stop()

    def test_upload_endpoint_requires_auth(self):
        """Unauthenticated requests to upload endpoint return 401."""
        url = f"/api/threads/{self.thread1.id}/files/"
        response = self.client.post(url, {}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_upload_fixtures_all_types(self):
        """Upload and parse one fixture per type: pdf, docx, xlsx, csv, pptx, image, txt, py, json, zip."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        url = f"/api/threads/{self.thread1.id}/files/"

        fixtures = [
            "sample.pdf",
            "sample.docx",
            "sample.xlsx",
            "sample.csv",
            "sample.pptx",
            "sample.png",
            "sample.txt",
            "sample.py",
            "sample.json",
            "sample.zip",
        ]

        for fname in fixtures:
            fpath = self.sample_dir / fname
            self.assertTrue(fpath.exists(), f"Sample file {fname} must exist.")
            with open(fpath, "rb") as f:
                uploaded = SimpleUploadedFile(fname, f.read())

            res = self.client.post(url, {"files": uploaded}, format="multipart")
            self.assertEqual(res.status_code, status.HTTP_201_CREATED, f"Failed uploading {fname}: {res.data}")
            self.assertEqual(len(res.data), 1)
            record = res.data[0]
            self.assertEqual(record["status"], "ready", f"File {fname} parsing failed: {record.get('error')}")
            self.assertTrue(record["summary"], f"File {fname} should have a summary.")
            self.assertTrue(record["parsed_path"], f"File {fname} should have a parsed_path.")

            # Validate Image dimension storage
            if fname == "sample.png":
                parsed = storage.read_parsed(record["parsed_path"])
                self.assertIn("image_info", parsed)
                self.assertEqual(parsed["image_info"]["width"], 32)
                self.assertEqual(parsed["image_info"]["height"], 32)

    def test_multipart_many_files_per_request(self):
        """Upload multiple files in a single request."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        url = f"/api/threads/{self.thread1.id}/files/"

        f1 = SimpleUploadedFile("batch1.txt", b"First batch text content.")
        f2 = SimpleUploadedFile("batch2.json", b'{"key": "value"}')

        res = self.client.post(url, {"files": [f1, f2]}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(len(res.data), 2)
        self.assertTrue(all(item["status"] == "ready" for item in res.data))

    def test_corrupt_file_handling(self):
        """Corrupt files get status failed with a clear message."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        url = f"/api/threads/{self.thread1.id}/files/"

        corrupt_pdf = SimpleUploadedFile("corrupt.pdf", b"%PDF-1.4 corrupt junk data not a real pdf body")
        res = self.client.post(url, {"files": corrupt_pdf}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        record = res.data[0]
        self.assertEqual(record["status"], "failed")
        self.assertIn("Parsing failed", record["error"])

    def test_wrong_extension_handling(self):
        """Executable disguised with .pdf extension gets status failed."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        url = f"/api/threads/{self.thread1.id}/files/"

        # DOS / Windows executable header
        fake_pdf = SimpleUploadedFile("evil.pdf", b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff")
        res = self.client.post(url, {"files": fake_pdf}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        record = res.data[0]
        self.assertEqual(record["status"], "failed")
        self.assertTrue("MIME type" in record["error"] or "not permitted" in record["error"])

    def test_oversized_file_handling(self):
        """File exceeding MAX_FILE_SIZE_BYTES receives failed status."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        url = f"/api/threads/{self.thread1.id}/files/"

        with patch("filesvc.validators.MAX_FILE_SIZE_BYTES", 100):
            big_file = SimpleUploadedFile("too_big.txt", b"A" * 500)
            res = self.client.post(url, {"files": big_file}, format="multipart")
            self.assertEqual(res.status_code, status.HTTP_201_CREATED)
            record = res.data[0]
            self.assertEqual(record["status"], "failed")
            self.assertIn("exceeds limit", record["error"])

    def test_oversized_request_handling(self):
        """Total request size exceeding MAX_REQUEST_SIZE_BYTES returns 400 Bad Request."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        url = f"/api/threads/{self.thread1.id}/files/"

        with patch("filesvc.views.MAX_REQUEST_SIZE_BYTES", 200):
            f1 = SimpleUploadedFile("f1.txt", b"A" * 150)
            f2 = SimpleUploadedFile("f2.txt", b"B" * 150)
            res = self.client.post(url, {"files": [f1, f2]}, format="multipart")
            self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
            self.assertIn("Total request size", res.data["detail"])

    def test_zip_safety_traversal(self):
        """ZIP archives with path traversal are rejected with status failed."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        url = f"/api/threads/{self.thread1.id}/files/"

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("../../etc/passwd", "root:x:0:0:")
        buf.seek(0)

        traversal_zip = SimpleUploadedFile("traversal.zip", buf.read())
        res = self.client.post(url, {"files": traversal_zip}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        record = res.data[0]
        self.assertEqual(record["status"], "failed")
        self.assertIn("traversal", record["error"])

    def test_zip_safety_bomb_and_nested_archive(self):
        """ZIP bomb exceeding uncompressed limit and nested archives are rejected."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        url = f"/api/threads/{self.thread1.id}/files/"

        # Nested archive check
        buf_nested = io.BytesIO()
        with zipfile.ZipFile(buf_nested, "w") as zf:
            zf.writestr("inner.zip", b"PK\x05\x06" + b"\x00" * 18)
        buf_nested.seek(0)

        nested_zip = SimpleUploadedFile("nested.zip", buf_nested.read())
        res = self.client.post(url, {"files": nested_zip}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data[0]["status"], "failed")
        self.assertIn("nested archive", res.data[0]["error"].lower())

        # Uncompressed size limit check
        with patch("filesvc.validators.MAX_ZIP_UNCOMPRESSED_SIZE", 50):
            buf_big = io.BytesIO()
            with zipfile.ZipFile(buf_big, "w") as zf:
                zf.writestr("big.txt", b"X" * 200)
            buf_big.seek(0)

            bomb_zip = SimpleUploadedFile("bomb.zip", buf_big.read())
            res = self.client.post(url, {"files": bomb_zip}, format="multipart")
            self.assertEqual(res.status_code, status.HTTP_201_CREATED)
            self.assertEqual(res.data[0]["status"], "failed")
            self.assertIn("exceeds limit", res.data[0]["error"])

    def test_cross_user_access_returns_404(self):
        """Cross-user upload, list, and download return 404."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user2}")

        # User 2 attempts to upload to User 1's thread
        upload_url = f"/api/threads/{self.thread1.id}/files/"
        f = SimpleUploadedFile("test.txt", b"Hello")
        res = self.client.post(upload_url, {"files": f}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

        # User 2 attempts to list User 1's thread files
        res = self.client.get(upload_url)
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

        # User 1 creates a file
        uf = UploadedFile.objects.create(
            thread=self.thread1,
            name="private.txt",
            mime="text/plain",
            size=11,
            path=str(self.sample_dir / "sample.txt"),
            status="ready",
        )

        # User 2 attempts to download User 1's file
        download_url = f"/api/files/{uf.id}/download/"
        res = self.client.get(download_url)
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

    def test_download_endpoint_security_headers_and_sanitization(self):
        """Download endpoint sets nosniff, sanitizes filename, and never renders HTML/SVG inline."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")

        # Normal text download
        uf_text = UploadedFile.objects.create(
            thread=self.thread1,
            name="my\"document\r\n.txt",
            mime="text/plain",
            size=10,
            path=str(self.sample_dir / "sample.txt"),
            status="ready",
        )
        res = self.client.get(f"/api/files/{uf_text.id}/download/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.headers.get("X-Content-Type-Options"), "nosniff")
        disposition = res.headers.get("Content-Disposition", "")
        self.assertIn("attachment", disposition)
        self.assertNotIn("\r", disposition)
        self.assertNotIn("\n", disposition)
        self.assertNotIn('""', disposition)

        # HTML file must NEVER be served inline or as text/html
        html_path = self.sample_dir / "sample_html.html"
        with open(html_path, "w", encoding="utf-8") as f:
            f.write("<script>alert('xss')</script>")

        uf_html = UploadedFile.objects.create(
            thread=self.thread1,
            name="test.html",
            mime="text/html",
            size=30,
            path=str(html_path),
            status="ready",
        )
        res_html = self.client.get(f"/api/files/{uf_html.id}/download/")
        self.assertEqual(res_html.status_code, status.HTTP_200_OK)
        self.assertEqual(res_html.headers.get("Content-Type"), "application/octet-stream")
        self.assertIn("attachment", res_html.headers.get("Content-Disposition", ""))

        # SVG file must NEVER be served inline
        svg_path = self.sample_dir / "sample_svg.svg"
        with open(svg_path, "w", encoding="utf-8") as f:
            f.write("<svg onload='alert(1)'></svg>")

        uf_svg = UploadedFile.objects.create(
            thread=self.thread1,
            name="vector.svg",
            mime="image/svg+xml",
            size=30,
            path=str(svg_path),
            status="ready",
        )
        res_svg = self.client.get(f"/api/files/{uf_svg.id}/download/")
        self.assertEqual(res_svg.status_code, status.HTTP_200_OK)
        self.assertEqual(res_svg.headers.get("Content-Type"), "application/octet-stream")

    def test_tool_read_file(self):
        """Test tool read_file for spreadsheets, images, truncation, and cross-user verification."""
        # 1. Spreadsheet reading (schema + sample by default)
        uf_csv = UploadedFile.objects.create(
            thread=self.thread1,
            name="sample.csv",
            mime="text/csv",
            size=77,
            path=str(self.sample_dir / "sample.csv"),
            status="ready",
        )
        from filesvc.service import file_service
        file_service.parse_file_sync(uf_csv.id)
        uf_csv.refresh_from_db()

        config_user1 = {"configurable": {"thread_id": str(self.thread1.id), "user_id": str(self.user1.id)}}
        res_default = read_file.invoke({"file_id": str(uf_csv.id)}, config=config_user1)
        self.assertIn("Total Rows", res_default)
        self.assertIn("Schema", res_default)

        # 2. Spreadsheet row range request
        res_range = read_file.invoke({"file_id": str(uf_csv.id), "row_start": 1, "row_end": 2}, config=config_user1)
        self.assertIn("Tokyo", res_range)
        self.assertIn("Rows 1 to 2", res_range)

        # 3. Image dimensions and vision payload
        uf_img = UploadedFile.objects.create(
            thread=self.thread1,
            name="sample.png",
            mime="image/png",
            size=116,
            path=str(self.sample_dir / "sample.png"),
            status="ready",
        )
        file_service.parse_file_sync(uf_img.id)
        uf_img.refresh_from_db()

        res_img = read_file.invoke({"file_id": str(uf_img.id)}, config=config_user1)
        self.assertTrue("Traffic Sources" in res_img or "image" in res_img.lower() or "png" in res_img.lower())

        # 4. Bounded output with truncation
        large_txt_path = self.sample_dir / "large.txt"
        with open(large_txt_path, "w", encoding="utf-8") as f:
            f.write("Line of content.\n" * 500)

        uf_large = UploadedFile.objects.create(
            thread=self.thread1,
            name="large.txt",
            mime="text/plain",
            size=5000,
            path=str(large_txt_path),
            status="ready",
        )
        file_service.parse_file_sync(uf_large.id)
        uf_large.refresh_from_db()

        res_large = read_file.invoke({"file_id": str(uf_large.id)}, config=config_user1)
        self.assertIn("Truncated: showing chars 0-8000 of", res_large)

        # 5. Cross-user verification
        config_user2 = {"configurable": {"thread_id": str(self.thread2.id), "user_id": str(self.user2.id)}}
        res_cross = read_file.invoke({"file_id": str(uf_csv.id)}, config=config_user2)
        self.assertIn("Error:", res_cross)
        self.assertTrue("access denied" in res_cross.lower() or "does not belong" in res_cross.lower())

    def test_throttling_classes(self):
        """Verify DRF throttling classes for upload and chat are properly scoped."""
        upload_throttle = UploadRateThrottle()
        chat_throttle = ChatRateThrottle()
        self.assertEqual(upload_throttle.scope, "upload")
        self.assertEqual(chat_throttle.scope, "chat")

    def test_each_generator_opens_correctly_and_contains_spec_content(self):
        """Verify each generator creates a valid file that can be opened and contains spec content."""
        import csv
        import docx
        import openpyxl
        from PIL import Image
        from pptx import Presentation
        from pypdf import PdfReader

        gen_dir = self.sample_dir / "generated_test_output"
        gen_dir.mkdir(parents=True, exist_ok=True)

        # 1. DOCX
        docx_path = str(gen_dir / "test.docx")
        docx_spec = {
            "title": "Quarterly Performance",
            "sections": [{"heading": "Executive Summary", "content": "Revenue grew by 24 percent."}],
            "tables": [{"headers": ["Metric", "Value"], "rows": [["Q1", "$1.2M"], ["Q2", "$1.5M"]]}],
        }
        get_generator("docx").generate(docx_spec, docx_path)
        doc = docx.Document(docx_path)
        self.assertEqual(doc.paragraphs[0].text, "Quarterly Performance")
        self.assertIn("Revenue grew", doc.paragraphs[2].text)
        self.assertEqual(doc.tables[0].rows[1].cells[0].text, "Q1")

        # 2. PDF
        pdf_path = str(gen_dir / "test.pdf")
        pdf_spec = {
            "title": "Annual Audit Report",
            "sections": [{"heading": "Findings", "content": "All compliance criteria met."}],
        }
        get_generator("pdf").generate(pdf_spec, pdf_path)
        reader = PdfReader(pdf_path)
        self.assertGreaterEqual(len(reader.pages), 1)
        pdf_text = reader.pages[0].extract_text()
        self.assertIn("Annual Audit Report", pdf_text)
        self.assertIn("All compliance criteria met", pdf_text)

        # 3. XLSX
        xlsx_path = str(gen_dir / "test.xlsx")
        xlsx_spec = {
            "sheet_name": "Sales",
            "headers": ["Product", "Units", "Revenue"],
            "rows": [["Widget A", 150, 4500], ["Widget B", 80, 3200]],
        }
        get_generator("xlsx").generate(xlsx_spec, xlsx_path)
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb["Sales"]
        self.assertEqual(ws.cell(1, 1).value, "Product")
        self.assertEqual(ws.cell(2, 2).value, 150)

        # 4. CSV
        csv_path = str(gen_dir / "test.csv")
        csv_spec = {
            "headers": ["Department", "Budget"],
            "rows": [["Engineering", 850000], ["Marketing", 420000]],
        }
        get_generator("csv").generate(csv_spec, csv_path)
        with open(csv_path, "r", encoding="utf-8") as f:
            rows = list(csv.reader(f))
            self.assertEqual(rows[0], ["Department", "Budget"])
            self.assertEqual(rows[1], ["Engineering", "850000"])

        # 5. PPTX
        pptx_path = str(gen_dir / "test.pptx")
        pptx_spec = {
            "title": "Keynote Presentation",
            "slides": [{"title": "Vision", "bullets": ["Customer Obsession", "Velocity"]}],
        }
        get_generator("pptx").generate(pptx_spec, pptx_path)
        prs = Presentation(pptx_path)
        self.assertEqual(prs.slides[0].shapes.title.text, "Keynote Presentation")
        self.assertEqual(prs.slides[1].shapes.title.text, "Vision")

        # 6. PNG Chart
        chart_path = str(gen_dir / "test_chart.png")
        chart_spec = {
            "chart_type": "bar",
            "title": "Quarterly Revenue",
            "x": ["Q1", "Q2", "Q3"],
            "y": [12, 19, 15],
            "x_label": "Quarter",
            "y_label": "Millions ($)",
        }
        get_generator("chart").generate(chart_spec, chart_path)
        with Image.open(chart_path) as img:
            self.assertEqual(img.format, "PNG")
            self.assertGreater(img.size[0], 200)

        # 7. MD
        md_path = str(gen_dir / "test.md")
        md_spec = {
            "title": "Architecture Overview",
            "sections": [{"heading": "Components", "bullets": ["Frontend", "Backend"]}],
        }
        get_generator("md").generate(md_spec, md_path)
        with open(md_path, "r", encoding="utf-8") as f:
            content = f.read()
            self.assertIn("# Architecture Overview", content)
            self.assertIn("- Frontend", content)

        # 8. Code
        code_path = str(gen_dir / "test.py")
        code_spec = {
            "language": "python",
            "code": "def calculate(a, b):\n    return a * b + 1\n",
        }
        get_generator("code").generate(code_spec, code_path)
        with open(code_path, "r", encoding="utf-8") as f:
            code_text = f.read()
            self.assertIn("return a * b + 1", code_text)

    def test_generate_file_tool_and_revision_parent_tracking(self):
        """Tool generate_file creates version 1, read_generated_spec inspects it, and based_on creates version 2 with parent."""
        config = {"configurable": {"thread_id": str(self.thread1.id), "user_id": str(self.user1.id)}}

        # Version 1 creation
        spec_v1 = {"title": "Release Plan v1", "content": "Initial release schedule."}
        res_v1 = generate_file.invoke(
            {"kind": "md", "name": "release_plan.md", "spec": spec_v1},
            config=config,
        )
        self.assertIn("Successfully generated md file", res_v1)
        self.assertIn("version: 1", res_v1)

        gf_v1 = GeneratedFile.objects.get(thread=self.thread1, name="release_plan.md", version=1)
        self.assertIsNone(gf_v1.parent)

        # Inspect spec using read_generated_spec
        read_res = read_generated_spec.invoke({"file_id": gf_v1.id}, config=config)
        self.assertIn("Release Plan v1", read_res)

        # Revision creation: version 2 based on v1
        spec_v2 = {"title": "Release Plan v2", "content": "Updated schedule with expedited release."}
        res_v2 = generate_file.invoke(
            {"kind": "md", "name": "release_plan.md", "spec": spec_v2, "based_on": gf_v1.id},
            config=config,
        )
        self.assertIn("Successfully generated md file", res_v2)
        self.assertIn("version: 2", res_v2)

        gf_v2 = GeneratedFile.objects.get(thread=self.thread1, name="release_plan.md", version=2)
        self.assertEqual(gf_v2.parent, gf_v1)
        self.assertEqual(gf_v2.version, 2)
        self.assertEqual(gf_v2.spec["title"], "Release Plan v2")

    def test_generator_failure_surfaces_to_model(self):
        """Invalid spec or missing parent surfaces clear error string to model, not silently dropped."""
        config = {"configurable": {"thread_id": str(self.thread1.id), "user_id": str(self.user1.id)}}

        # Invalid chart type
        bad_chart_spec = {"chart_type": "3d_surface_hologram", "x": [1, 2], "y": [3, 4]}
        res_fail = generate_file.invoke(
            {"kind": "chart", "name": "bad.png", "spec": bad_chart_spec},
            config=config,
        )
        self.assertIn("Error generating file:", res_fail)
        self.assertIn("Unsupported chart_type", res_fail)

        # Unknown kind
        res_unknown = generate_file.invoke(
            {"kind": "unknown_file_type", "name": "bad.xyz", "spec": {"data": 123}},
            config=config,
        )
        self.assertIn("Error generating file:", res_unknown)
        self.assertIn("Unknown generator kind", res_unknown)

        # Nonexistent parent
        res_missing_parent = generate_file.invoke(
            {"kind": "md", "name": "rev.md", "spec": {"title": "T"}, "based_on": 999999},
            config=config,
        )
        self.assertIn("Error:", res_missing_parent)
        self.assertIn("Parent generated file id '999999' not found", res_missing_parent)

    def test_manifest_lists_generated_files_with_spec_summary(self):
        """File manifest lists generated files with ids and a short spec summary."""
        from agent.nodes import _fetch_file_manifest
        from asgiref.sync import async_to_sync

        spec = {
            "title": "System Architecture",
            "sections": [{"heading": "Microservices", "content": "Documented"}],
        }
        gf = GeneratedFile.objects.create(
            thread=self.thread1,
            name="architecture.docx",
            kind="docx",
            version=1,
            spec=spec,
            path="simulated_path.docx",
        )

        manifest = async_to_sync(_fetch_file_manifest)(str(self.thread1.id))
        gf_entry = next((item for item in manifest if item["id"] == str(gf.id)), None)
        self.assertIsNotNone(gf_entry)
        self.assertEqual(gf_entry["name"], "architecture.docx")
        self.assertEqual(gf_entry["kind"], "docx")
        self.assertIn("DOCX", gf_entry["spec_summary"])
        self.assertIn("System Architecture", gf_entry["spec_summary"])

    def test_generated_file_download_endpoint_and_security(self):
        """Authenticated download endpoint works for generated files with security headers and cross-user 404."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")

        temp_gen_path = self.sample_dir / "download_test.md"
        with open(temp_gen_path, "w", encoding="utf-8") as f:
            f.write("# Generated Content for Download")

        gf = GeneratedFile.objects.create(
            thread=self.thread1,
            name="my\"report\r\n.md",
            kind="md",
            version=1,
            spec={"title": "Test"},
            path=str(temp_gen_path),
        )

        res = self.client.get(f"/api/generated/{gf.id}/download/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.headers.get("X-Content-Type-Options"), "nosniff")
        disposition = res.headers.get("Content-Disposition", "")
        self.assertIn("attachment", disposition)
        self.assertNotIn("\r", disposition)
        self.assertNotIn("\n", disposition)

        # Cross-user access returns 404
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user2}")
        res_cross = self.client.get(f"/api/generated/{gf.id}/download/")
        self.assertEqual(res_cross.status_code, status.HTTP_404_NOT_FOUND)

        # Generated HTML file must NEVER be rendered inline
        html_gen_path = self.sample_dir / "gen_html.html"
        with open(html_gen_path, "w", encoding="utf-8") as f:
            f.write("<html><body>Generated HTML</body></html>")

        gf_html = GeneratedFile.objects.create(
            thread=self.thread1,
            name="page.html",
            kind="code",
            version=1,
            spec={"code": "<html></html>"},
            path=str(html_gen_path),
        )
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        res_html = self.client.get(f"/api/generated/{gf_html.id}/download/")
        self.assertEqual(res_html.status_code, status.HTTP_200_OK)
        self.assertEqual(res_html.headers.get("Content-Type"), "application/octet-stream")

    async def test_fake_llm_and_sse_events_during_file_generation(self):
        """FAKE_LLM drives file generation without a key, emitting status during tool calls and a file event when done."""
        from chat.streaming import stream_chat_turn

        events = []
        async for chunk in stream_chat_turn(thread=self.thread1, user_prompt="Please generate a project report"):
            events.append(chunk)

        full_output = "".join(events)
        # Verify status event emitted during tool call
        self.assertIn("event: status", full_output)
        self.assertIn("calling tool generate_file", full_output)

        # Verify file event emitted when done with id, name, kind, version, download_url
        self.assertIn("event: file", full_output)
        self.assertIn("/api/generated/", full_output)
        self.assertIn("event: done", full_output)

        # Verify GeneratedFile row created in DB
        gf = await GeneratedFile.objects.filter(thread=self.thread1, name="report.md").afirst()
        self.assertIsNotNone(gf)
        self.assertEqual(gf.version, 1)

    def test_xlsx_generation_and_openpyxl_cached_values(self):
        """XLSX generator writes computed values and number formats so openpyxl data_only=True reads them."""
        import openpyxl
        import tempfile
        from filesvc.generators.registry import get_generator

        gen = get_generator("xlsx")
        with tempfile.TemporaryDirectory() as tmp_dir:
            out_path = os.path.join(tmp_dir, "test_cached_calc.xlsx")
            spec = {
                "headers": ["Month", "Revenue (INR)", "MoM Growth"],
                "rows": [
                    ["Jan", {"value": 1000, "number_format": "#,##0"}, ""],
                    ["Feb", {"value": 1500, "number_format": "#,##0"}, {"value": 0.5, "number_format": "0.0%"}],
                    ["Total", {"value": 2500, "formula": "=SUM(B2:B3)", "bold": True, "number_format": "#,##0"}, ""],
                ],
            }
            gen.generate(spec, out_path)

            # Verify with openpyxl data_only=True
            wb_cached = openpyxl.load_workbook(out_path, data_only=True)
            sheet = wb_cached.active
            self.assertEqual(sheet.max_row, 4)
            self.assertEqual(sheet.max_column, 3)

            # Check total cell (B4) has cached value 2500 and is not None
            tot_val = sheet.cell(4, 2).value
            self.assertEqual(tot_val, 2500)

            # Check first period growth (C2) is None or empty string
            self.assertIn(sheet.cell(2, 3).value, (None, ""))

            # Check total row growth (C4) is None or empty string
            self.assertIn(sheet.cell(4, 3).value, (None, ""))

    def test_xlsx_verification_catches_sum_mismatch_and_non_blank_growth(self):
        """Verification raises ValueError if total sum mismatches or growth column rules are violated."""
        import tempfile
        from filesvc.generators.verification import verify_generated_file
        from filesvc.generators.registry import get_generator

        gen = get_generator("xlsx")
        with tempfile.TemporaryDirectory() as tmp_dir:
            bad_sum_path = os.path.join(tmp_dir, "bad_sum.xlsx")
            spec_bad_sum = {
                "headers": ["Month", "Revenue (INR)"],
                "rows": [
                    ["Jan", 1000],
                    ["Feb", 1500],
                    ["Total", 9999],  # Incorrect sum
                ],
            }
            gen.generate(spec_bad_sum, bad_sum_path)
            with self.assertRaises(ValueError) as ctx:
                verify_generated_file("xlsx", bad_sum_path, spec_bad_sum)
            self.assertIn("does not match sum of data rows", str(ctx.exception))

            # Check invalid first period growth raises ValueError
            bad_growth_path = os.path.join(tmp_dir, "bad_growth.xlsx")
            spec_bad_growth = {
                "headers": ["Month", "Revenue (INR)", "MoM Growth"],
                "rows": [
                    ["Jan", 1000, 0.10],  # Must be blank
                    ["Feb", 1500, 0.50],
                    ["Total", 2500, ""],
                ],
            }
            gen.generate(spec_bad_growth, bad_growth_path)
            with self.assertRaises(ValueError) as ctx2:
                verify_generated_file("xlsx", bad_growth_path, spec_bad_growth)
            self.assertIn("first period growth cell C2 must be blank", str(ctx2.exception))

    def test_generated_file_preview_endpoint(self):
        """Preview endpoint returns structured table with cached values and handles cross-tenant permissions."""
        import tempfile
        from filesvc.generators.registry import get_generator

        gen = get_generator("xlsx")
        tmp_dir = tempfile.mkdtemp()
        self.addCleanup(lambda: shutil.rmtree(tmp_dir, ignore_errors=True))
        import shutil
        xlsx_path = os.path.join(tmp_dir, "preview_sample.xlsx")
        spec = {
            "headers": ["Month", "Revenue (INR)"],
            "rows": [
                ["Jan", {"value": 1240000, "number_format": "#,##0"}],
                ["Feb", {"value": 1085000, "number_format": "#,##0"}],
                ["Total", {"value": 2325000, "bold": True, "number_format": "#,##0"}],
            ],
        }
        gen.generate(spec, xlsx_path)

        gf = GeneratedFile.objects.create(
            thread=self.thread1,
            name="preview_revenue.xlsx",
            kind="xlsx",
            version=1,
            spec=spec,
            path=xlsx_path,
        )

        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user1}")
        res = self.client.get(f"/api/generated/{gf.id}/preview/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        self.assertEqual(data["type"], "table")
        self.assertEqual(data["headers"], ["Month", "Revenue (INR)"])
        self.assertEqual(len(data["rows"]), 3)
        self.assertEqual(data["rows"][2], ["Total", "2,325,000"])
        self.assertFalse(data["has_uncalculated_formulas"])
        self.assertIsNone(data["warning"])

        # Cross-user unauthorized returns 404
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.token_user2}")
        res_cross = self.client.get(f"/api/generated/{gf.id}/preview/")
        self.assertEqual(res_cross.status_code, status.HTTP_404_NOT_FOUND)

