import docx
from filesvc.parsers.base import BaseParser, ParsedResult

class DocxParser(BaseParser):
    """Parser for Microsoft Word (.docx) files using python-docx."""

    def parse(self, file_path: str, original_filename: str) -> ParsedResult:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        table_rows = []
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                table_rows.append(" | ".join(cells))

        parts = paragraphs + (["--- Tables ---"] + table_rows if table_rows else [])
        full_text = "\n".join(parts).strip()
        preview = (full_text[:200].replace("\n", " ").strip()) if full_text else "Empty document"
        summary = f"Word document '{original_filename}': {len(paragraphs)} paragraph(s), {len(doc.tables)} table(s). Preview: {preview}"
        return ParsedResult(
            text=full_text,
            summary=summary,
            metadata={
                "paragraph_count": len(paragraphs),
                "table_count": len(doc.tables),
            },
        )
