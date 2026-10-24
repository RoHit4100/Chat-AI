from pptx import Presentation
from filesvc.parsers.base import BaseParser, ParsedResult

class PptxParser(BaseParser):
    """Parser for PowerPoint (.pptx) presentations using python-pptx."""

    def parse(self, file_path: str, original_filename: str) -> ParsedResult:
        prs = Presentation(file_path)
        slides_text = []
        for i, slide in enumerate(prs.slides):
            slide_lines = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_lines.append(text)
            if slide_lines:
                slides_text.append(f"--- Slide {i + 1} ---\n" + "\n".join(slide_lines))

        full_text = "\n\n".join(slides_text).strip()
        preview = (full_text[:200].replace("\n", " ").strip()) if full_text else "Empty presentation"
        summary = f"PowerPoint presentation '{original_filename}': {len(prs.slides)} slide(s). Preview: {preview}"
        return ParsedResult(
            text=full_text,
            summary=summary,
            metadata={"slide_count": len(prs.slides)},
        )
