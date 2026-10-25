import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from filesvc.generators.base import BaseGenerator
from filesvc.generators.specs import validate_spec_for_kind

class DocxGenerator(BaseGenerator):
    """Generator for Microsoft Word (.docx) documents with rich styling and ATS resume layout."""
    default_extension = ".docx"

    def validate_spec(self, spec: dict) -> None:
        super().validate_spec(spec)
        validate_spec_for_kind("docx", spec)

    def generate(self, spec: dict, output_path: str) -> None:
        self.validate_spec(spec)
        doc = docx.Document()

        # Set page margins (0.5 inch default for compact layout)
        sections_page = doc.sections
        for s in sections_page:
            margins = spec.get("margins")
            if isinstance(margins, dict):
                s.top_margin = Pt(margins.get("top", 36))
                s.bottom_margin = Pt(margins.get("bottom", 36))
                s.left_margin = Pt(margins.get("left", 36))
                s.right_margin = Pt(margins.get("right", 36))
            elif isinstance(margins, (int, float)):
                s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Pt(margins)
            else:
                s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(0.5)

        # Title / Name
        title = spec.get("title")
        if title:
            h = doc.add_paragraph()
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = h.add_run(str(title))
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            h.paragraph_format.space_after = Pt(2)

        # Subtitle / Role
        subtitle = spec.get("subtitle")
        if subtitle:
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = sub.add_run(str(subtitle))
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            sub.paragraph_format.space_after = Pt(3)

        # Header / Contact lines
        header = spec.get("header") or spec.get("contact")
        if header:
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if isinstance(header, list):
                text = "  •  ".join(str(item).strip() for item in header if str(item).strip())
            elif isinstance(header, dict):
                text = "  •  ".join(f"{k}: {v}" for k, v in header.items() if v)
            else:
                text = str(header)
            run = c.add_run(text)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(71, 85, 105)
            c.paragraph_format.space_after = Pt(6)

        # Description / Author if present
        description = spec.get("description")
        if description:
            p = doc.add_paragraph(str(description))
            p.paragraph_format.space_after = Pt(4)

        # Top-level simple paragraphs
        for p_text in spec.get("paragraphs", []):
            p = doc.add_paragraph(str(p_text))
            p.paragraph_format.space_after = Pt(4)

        # Render sections
        for sec in spec.get("sections", []):
            if isinstance(sec, str):
                p = doc.add_paragraph(sec)
                p.paragraph_format.space_after = Pt(4)
                continue
            if not isinstance(sec, dict):
                continue

            heading = sec.get("heading") or sec.get("title")
            if heading:
                h = doc.add_paragraph()
                h.paragraph_format.space_before = Pt(8)
                h.paragraph_format.space_after = Pt(2)
                h.paragraph_format.keep_with_next = True
                run = h.add_run(str(heading).upper())
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(15, 23, 42)

            # Section content / paragraph
            content = sec.get("content") or sec.get("paragraph")
            if content:
                p = doc.add_paragraph(str(content))
                p.paragraph_format.space_after = Pt(3)

            # Section paragraphs
            for p_text in sec.get("paragraphs", []):
                p = doc.add_paragraph(str(p_text))
                p.paragraph_format.space_after = Pt(3)

            # Groups (e.g. categorized skills)
            for grp in sec.get("groups", []):
                if isinstance(grp, dict):
                    label = grp.get("label") or grp.get("name") or ""
                    items = grp.get("items") or grp.get("value") or grp.get("skills") or ""
                    if isinstance(items, list):
                        items_str = ", ".join(str(i) for i in items)
                    else:
                        items_str = str(items)
                    gp = doc.add_paragraph()
                    gp.paragraph_format.space_after = Pt(2)
                    if label:
                        r_lbl = gp.add_run(f"{label}: " if not label.endswith(":") else f"{label} ")
                        r_lbl.bold = True
                    gp.add_run(items_str)
                elif isinstance(grp, str):
                    p = doc.add_paragraph(grp)
                    p.paragraph_format.space_after = Pt(2)

            # Detailed entries (Work Experience, Projects, Education)
            for entry in sec.get("entries", []):
                if isinstance(entry, dict):
                    e_title = entry.get("title", "")
                    e_sub = entry.get("subtitle") or entry.get("company") or entry.get("organization") or ""
                    e_date = entry.get("date") or entry.get("dates") or ""
                    e_loc = entry.get("location", "")
                    e_desc = entry.get("description") or entry.get("content") or ""
                    e_bullets = entry.get("bullets", [])

                    # Row 1: Title and Date
                    ep = doc.add_paragraph()
                    ep.paragraph_format.space_before = Pt(3)
                    ep.paragraph_format.space_after = Pt(1)
                    ep.paragraph_format.keep_with_next = True
                    r_title = ep.add_run(str(e_title))
                    r_title.bold = True
                    if e_date:
                        ep.add_run(f"  |  {e_date}")

                    # Row 2: Subtitle / Organization & Location
                    sub_parts = [p for p in [e_sub, e_loc] if p]
                    if sub_parts:
                        sp = doc.add_paragraph()
                        sp.paragraph_format.space_after = Pt(2)
                        r_sub = sp.add_run("  |  ".join(sub_parts))
                        r_sub.italic = True
                        r_sub.font.color.rgb = RGBColor(51, 65, 85)

                    # Row 3: Description if present
                    if e_desc:
                        dp = doc.add_paragraph(str(e_desc))
                        dp.paragraph_format.space_after = Pt(2)

                    # Bullets
                    for b in e_bullets:
                        bp = doc.add_paragraph(str(b), style="List Bullet")
                        bp.paragraph_format.space_after = Pt(1.5)
                elif isinstance(entry, str):
                    p = doc.add_paragraph(entry)
                    p.paragraph_format.space_after = Pt(2)

            # Section top-level bullets
            for bullet in sec.get("bullets", []):
                bp = doc.add_paragraph(str(bullet), style="List Bullet")
                bp.paragraph_format.space_after = Pt(1.5)

            # Key-values
            kvs = sec.get("key_values", [])
            if isinstance(kvs, dict):
                kvs = [{k: v} for k, v in kvs.items()]
            for kv in kvs:
                if isinstance(kv, dict):
                    for k, v in kv.items():
                        kvp = doc.add_paragraph()
                        kvp.paragraph_format.space_after = Pt(2)
                        r_k = kvp.add_run(f"{k}: ")
                        r_k.bold = True
                        kvp.add_run(str(v))

            # Tables in section
            for tbl in sec.get("tables", []):
                self._render_table(doc, tbl)

        # Top-level tables
        for tbl in spec.get("tables", []):
            self._render_table(doc, tbl)

        doc.save(output_path)

    def _render_table(self, doc: docx.Document, tbl: dict) -> None:
        """Helper to render tables in DOCX."""
        headers = tbl.get("headers", [])
        rows = tbl.get("rows", [])
        if not headers and not rows:
            return
        col_count = len(headers) if headers else len(rows[0])
        table = doc.add_table(rows=0, cols=col_count)
        table.style = "Table Grid"

        if headers:
            hdr_cells = table.add_row().cells
            for i, h in enumerate(headers[:col_count]):
                hdr_cells[i].text = str(h)
                for p in hdr_cells[i].paragraphs:
                    for r in p.runs:
                        r.bold = True

        for r in rows:
            row_cells = table.add_row().cells
            for i, cell_val in enumerate(r[:col_count]):
                row_cells[i].text = str(cell_val)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)

    def summarize_spec(self, spec: dict) -> str:
        title = spec.get("title") or "Document"
        sections = len(spec.get("sections", [])) + len(spec.get("paragraphs", []))
        tables = len(spec.get("tables", []))
        return f"DOCX '{title}' ({sections} section(s), {tables} table(s))"
