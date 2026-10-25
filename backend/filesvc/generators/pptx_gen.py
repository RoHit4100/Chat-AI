from pptx import Presentation
from pptx.util import Inches, Pt
from filesvc.generators.base import BaseGenerator
from filesvc.generators.specs import validate_spec_for_kind

class PptxGenerator(BaseGenerator):
    """Generator for PowerPoint (.pptx) presentations using python-pptx."""
    default_extension = ".pptx"

    def validate_spec(self, spec: dict) -> None:
        super().validate_spec(spec)
        validate_spec_for_kind("pptx", spec)

    def generate(self, spec: dict, output_path: str) -> None:
        self.validate_spec(spec)
        prs = Presentation()

        title = spec.get("title")
        subtitle = spec.get("subtitle")

        # Slide 1: Title slide layout (layout 0)
        if title:
            title_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(title_layout)
            slide.shapes.title.text = str(title)
            if subtitle and len(slide.placeholders) > 1:
                slide.placeholders[1].text = str(subtitle)

        # Remaining slides: Title and content layout (layout 1)
        for s_spec in spec.get("slides", []):
            if isinstance(s_spec, dict):
                slide_layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(slide_layout)
                stitle = s_spec.get("title", "")
                if stitle:
                    slide.shapes.title.text = str(stitle)

                content_box = slide.placeholders[1].text_frame
                content_text = s_spec.get("content")
                bullets = s_spec.get("bullets", [])

                if content_text:
                    content_box.text = str(content_text)

                for b in bullets:
                    p = content_box.add_paragraph()
                    p.text = str(b)
                    p.level = 0
            elif isinstance(s_spec, str):
                slide_layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(slide_layout)
                slide.shapes.title.text = s_spec

        prs.save(output_path)

    def summarize_spec(self, spec: dict) -> str:
        title = spec.get("title") or "Presentation"
        slides = len(spec.get("slides", [])) + (1 if spec.get("title") else 0)
        return f"PPTX '{title}' ({slides} slide(s))"
