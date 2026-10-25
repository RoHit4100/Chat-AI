import os
import shutil
import tempfile
import subprocess
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from filesvc.generators.base import BaseGenerator
from filesvc.generators.specs import validate_spec_for_kind

def recalculate_with_libreoffice(file_path: str) -> bool:
    """Recalculate formulas in Excel workbook using headless LibreOffice Calc to write cached values."""
    abs_path = os.path.abspath(file_path)
    if not os.path.exists(abs_path):
        return False

    out_dir = tempfile.mkdtemp(prefix="soffice_out_")
    profile_dir = tempfile.mkdtemp(prefix="soffice_profile_")
    try:
        env = os.environ.copy()
        env["SAL_USE_VCLPLUGIN"] = "svp"
        cmd = [
            "soffice",
            f"-env:UserInstallation=file://{profile_dir}",
            "--headless",
            "--convert-to", "xlsx",
            "--outdir", out_dir,
            abs_path,
        ]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=30, env=env)
        base_name = os.path.basename(abs_path)
        conv_path = os.path.join(out_dir, base_name)
        if os.path.exists(conv_path) and os.path.getsize(conv_path) > 0:
            shutil.copy2(conv_path, abs_path)
            return True
        return False
    except Exception:
        return False
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)
        shutil.rmtree(profile_dir, ignore_errors=True)

def _to_float(v: Any) -> float | None:
    if isinstance(v, dict):
        v = v.get("value")
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "").replace("₹", "").replace("$", ""))
    except (ValueError, TypeError):
        return None

class XlsxGenerator(BaseGenerator):
    """Generator for Excel (.xlsx) workbooks from structured spec with computed values and formatting."""
    default_extension = ".xlsx"

    def validate_spec(self, spec: dict) -> None:
        super().validate_spec(spec)
        validate_spec_for_kind("xlsx", spec)

    def generate(self, spec: dict, output_path: str) -> None:
        self.validate_spec(spec)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default active sheet

        sheets_spec = spec.get("sheets")
        if not sheets_spec:
            sheets_spec = [{
                "name": spec.get("sheet_name", "Data"),
                "headers": spec.get("headers", []),
                "rows": spec.get("rows", []),
                "total_row": spec.get("total_row"),
                "computed_columns": spec.get("computed_columns", []),
            }]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        total_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True)
        total_border = Border(
            top=Side(style="thin", color="808080"),
            bottom=Side(style="double", color="000000"),
        )
        data_font = Font(name="Calibri", size=11)

        has_formulas = False

        for s_spec in sheets_spec:
            ws = wb.create_sheet(title=str(s_spec.get("name", "Sheet"))[:31])
            headers = [str(h) for h in (s_spec.get("headers") or spec.get("headers") or [])]
            raw_rows = list(s_spec.get("rows") or spec.get("rows") or [])
            total_row_cfg = s_spec.get("total_row") or spec.get("total_row")
            computed_cols_cfg = s_spec.get("computed_columns") or spec.get("computed_columns") or []

            # Separate existing total row if present in raw_rows
            data_rows = []
            has_total_in_rows = False
            for r in raw_rows:
                r_list = list(r) if isinstance(r, (list, tuple)) else [r]
                first_val = r_list[0] if r_list else None
                if isinstance(first_val, dict):
                    first_val = first_val.get("value")
                if first_val and str(first_val).strip().lower() == "total":
                    has_total_in_rows = True
                else:
                    data_rows.append(r_list)

            # Process computed_columns (e.g. MoM growth %)
            for comp in computed_cols_cfg:
                comp_header = comp.get("header") or "MoM growth %"
                source_col = comp.get("source") or (headers[1] if len(headers) > 1 else None)
                kind = comp.get("kind", "pct_change")

                if comp_header not in headers:
                    headers.append(comp_header)
                    comp_col_idx = len(headers) - 1
                else:
                    comp_col_idx = headers.index(comp_header)

                source_col_idx = headers.index(source_col) if (source_col and source_col in headers) else None

                if kind == "pct_change" and source_col_idx is not None:
                    for i, drow in enumerate(data_rows):
                        while len(drow) <= comp_col_idx:
                            drow.append(None)
                        if i == 0:
                            # First data row stays blank in growth column
                            drow[comp_col_idx] = {"value": "", "number_format": "0.0%"}
                        else:
                            prev_num = _to_float(data_rows[i - 1][source_col_idx])
                            curr_num = _to_float(drow[source_col_idx])
                            if prev_num is not None and curr_num is not None and prev_num != 0:
                                pct = round((curr_num - prev_num) / prev_num, 6)
                                drow[comp_col_idx] = {"value": pct, "number_format": "0.0%"}
                            else:
                                drow[comp_col_idx] = {"value": "", "number_format": "0.0%"}

            # Process total_row (compute totals in Python)
            total_row_cells = None
            if total_row_cfg or has_total_in_rows:
                label = "Total"
                sum_cols = []
                if isinstance(total_row_cfg, dict):
                    label = total_row_cfg.get("label", "Total") or "Total"
                    sum_cols = total_row_cfg.get("columns", [])

                total_row_cells = []
                for c_idx, h in enumerate(headers):
                    h_lower = h.lower()
                    if c_idx == 0:
                        total_row_cells.append({"value": label, "bold": True})
                    elif any(comp.get("header") == h for comp in computed_cols_cfg) or "growth" in h_lower or "%" in h_lower:
                        # Total row stays blank in growth column
                        total_row_cells.append({"value": "", "bold": True})
                    elif (sum_cols and h in sum_cols) or (not sum_cols and any(k in h_lower for k in ["revenue", "sales", "amount", "inr", "total", "units"])):
                        # Sum data rows in Python (including January!)
                        col_sum = 0.0
                        has_num = False
                        for drow in data_rows:
                            if c_idx < len(drow):
                                val_num = _to_float(drow[c_idx])
                                if val_num is not None:
                                    col_sum += val_num
                                    has_num = True
                        if has_num:
                            final_sum = int(col_sum) if col_sum.is_integer() else col_sum
                            total_row_cells.append({"value": final_sum, "number_format": "#,##0", "bold": True})
                        else:
                            total_row_cells.append({"value": "", "bold": True})
                    else:
                        total_row_cells.append({"value": "", "bold": True})

            final_rows = data_rows + ([total_row_cells] if total_row_cells else [])

            # Write header row (row 1, bold)
            current_row = 1
            if headers:
                for col_idx, h in enumerate(headers, start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=str(h) if h is not None else "")
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                current_row += 1

            # Write data and total rows
            for row in final_rows:
                if not isinstance(row, (list, tuple)):
                    row = [row]

                is_total_row = False
                if len(row) > 0:
                    first_val = row[0]
                    if isinstance(first_val, dict):
                        first_val = first_val.get("value")
                    if first_val and str(first_val).strip().lower() == "total":
                        is_total_row = True

                for col_idx, item in enumerate(row, start=1):
                    cell = ws.cell(row=current_row, column=col_idx)

                    val = None
                    formula = None
                    num_fmt = None
                    bold = is_total_row

                    if isinstance(item, dict):
                        val = item.get("value")
                        formula = item.get("formula")
                        num_fmt = item.get("number_format")
                        if item.get("bold") is not None:
                            bold = bool(item.get("bold"))
                    else:
                        if isinstance(item, str) and item.startswith("="):
                            formula = item
                        else:
                            val = item

                    if formula:
                        cell.value = str(formula)
                        has_formulas = True
                    elif val is not None and val != "":
                        cell.value = val
                    else:
                        cell.value = None

                    # Infer / apply number formatting
                    header_text = str(headers[col_idx - 1] if col_idx <= len(headers) else "").lower()
                    if num_fmt:
                        cell.number_format = str(num_fmt)
                    elif isinstance(val, float) and ("growth" in header_text or "%" in header_text):
                        cell.number_format = "0.0%"
                    elif isinstance(val, (int, float)) and any(k in header_text for k in ["revenue", "amount", "total", "sales", "units", "inr"]):
                        cell.number_format = "#,##0"

                    if is_total_row:
                        cell.font = total_font
                        cell.fill = total_fill
                        cell.border = total_border
                    elif bold:
                        cell.font = Font(name="Calibri", size=11, bold=True)
                    else:
                        cell.font = data_font

                current_row += 1

            # Auto-fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    v_str = str(cell.value or "")
                    if cell.number_format and "0" in cell.number_format and len(v_str) < 10:
                        v_str = v_str + "    "
                    max_len = max(max_len, len(v_str))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        wb.save(output_path)

        if has_formulas:
            recalculate_with_libreoffice(output_path)

    def summarize_spec(self, spec: dict) -> str:
        sheets = spec.get("sheets")
        if sheets:
            count = len(sheets)
            total_rows = sum(len(s.get("rows", [])) for s in sheets)
            return f"XLSX with {count} sheet(s), {total_rows} row(s)"
        rows = len(spec.get("rows", []))
        cols = len(spec.get("headers", []))
        return f"XLSX with {rows} row(s), {cols} column(s)"
