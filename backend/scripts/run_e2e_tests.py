import os
import sys
import asyncio
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
os.environ["DJANGO_ALLOW_ASYNC_UNSAFE"] = "true"
import django
django.setup()

from datetime import datetime, timezone, timedelta
from unittest.mock import patch, MagicMock
from cryptography.hazmat.primitives.asymmetric import rsa
import jwt

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient
from rest_framework import status

from chat.models import Thread, Message, UploadedFile, GeneratedFile
from chat.streaming import stream_chat_turn
from filesvc.parsers.registry import get_parser_for_file
from filesvc.generators.registry import get_generator

User = get_user_model()
SAMPLE_DIR = BASE_DIR / "sample_files"

def make_auth_client(user):
    test_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    mock_jwks = MagicMock()
    mock_jwks.get_signing_key_from_jwt.return_value = MagicMock(key=test_key.public_key())
    
    token = jwt.encode(
        {
            "sub": str(user.username),
            "email": str(user.email),
            "iss": settings.KEYCLOAK_PUBLIC_URL,
            "azp": settings.KEYCLOAK_CLIENT_ID,
            "iat": int(datetime.now(timezone.utc).timestamp()),
            "exp": int((datetime.now(timezone.utc) + timedelta(hours=2)).timestamp()),
        },
        test_key,
        algorithm="RS256",
        headers={"kid": "e2e-kid"}
    )
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")
    return client, mock_jwks

async def run_chat(thread, prompt: str) -> tuple[str, list[dict]]:
    """Helper to run stream_chat_turn and collect tokens, status, files."""
    tokens = []
    files = []
    statuses = []
    
    async for event_str in stream_chat_turn(thread, user_prompt=prompt):
        lines = [l for l in event_str.strip().split("\n") if l]
        event_name = ""
        data = {}
        for line in lines:
            if line.startswith("event:"):
                event_name = line.split(":", 1)[1].strip()
            elif line.startswith("data:"):
                import json
                try:
                    data = json.loads(line.split(":", 1)[1].strip())
                except:
                    data = {}
        if event_name == "token":
            tokens.append(data.get("content", ""))
        elif event_name == "file":
            files.append(data)
        elif event_name == "status":
            statuses.append(data)
            
    full_text = "".join(tokens).strip()
    return full_text, files

async def test_file_reading_all_types(user, client, mock_jwks):
    """Test 1: Upload each allowed file type and ask a question answerable from its contents."""
    print("\n--- TEST 1: File reading for all allowed types ---")
    file_tests = [
        ("sample.pdf", "What is the document title in the uploaded PDF file?", "Sample PDF Title"),
        ("sample.docx", "What is the document heading in the uploaded Word document?", "Sample Document"),
        ("sample.xlsx", "What columns are present in the uploaded Excel spreadsheet?", "Score"),
        ("sample.csv", "What columns are in the uploaded CSV file?", "city"),
        ("sample.pptx", "What is the presentation slide title in the uploaded PPTX?", "Sample Presentation"),
        ("sample.png", "What are the dimensions and format of the uploaded image?", "32"),
        ("sample.txt", "What greeting text is inside the uploaded text file?", "Hello World"),
        ("sample.py", "What function is defined in the uploaded python file and what does it return?", "hello"),
        ("sample.json", "What is the value of the name field in the uploaded JSON file?", "Chat-AI"),
        ("sample.zip", "What filename is inside the uploaded ZIP archive?", "inner.txt"),
    ]
    
    with patch("chat.authentication.get_jwks_client", return_value=mock_jwks):
        for fname, prompt, expected in file_tests:
            thread = Thread.objects.create(user=user, title=f"Test {fname}")
            fpath = SAMPLE_DIR / fname
            with open(fpath, "rb") as f:
                upload_file = SimpleUploadedFile(fname, f.read())
                resp = client.post(f"/api/threads/{thread.id}/files/", {"files": upload_file}, format="multipart")
                assert resp.status_code == 201, f"Failed upload {fname}: {resp.data}"
                
            ans, _ = await run_chat(thread, prompt)
            print(f"[{fname}] Response: {ans[:100]}...")
            assert expected.lower() in ans.lower(), f"Expected '{expected}' in answer for {fname}, got: {ans}"
            print(f"[{fname}] PASS")

async def test_two_files_in_one_message(user, client, mock_jwks):
    """Test 2: Upload two files in one message and ask about both."""
    print("\n--- TEST 2: Two files in one message ---")
    thread = Thread.objects.create(user=user, title="Two Files Test")
    with patch("chat.authentication.get_jwks_client", return_value=mock_jwks):
        with open(SAMPLE_DIR / "sample.txt", "rb") as f1, open(SAMPLE_DIR / "sample.json", "rb") as f2:
            u1 = SimpleUploadedFile("sample.txt", f1.read())
            u2 = SimpleUploadedFile("sample.json", f2.read())
            resp = client.post(f"/api/threads/{thread.id}/files/", {"files": [u1, u2]}, format="multipart")
            assert resp.status_code == 201
            
        ans, _ = await run_chat(
            thread,
            "I uploaded two files: sample.txt and sample.json. What greeting is in the txt file, and what name is in the json file?"
        )
        print(f"Two files response: {ans}")
        assert "hello" in ans.lower()
        assert "chat-ai" in ans.lower()
        print("[Two files] PASS")

async def test_ask_15_messages_later(user, client, mock_jwks):
    """Test 3: Ask about the first file 15 messages later."""
    print("\n--- TEST 3: Ask about first file 15 messages later ---")
    thread = Thread.objects.create(user=user, title="15 Messages Later Test")
    with patch("chat.authentication.get_jwks_client", return_value=mock_jwks):
        with open(SAMPLE_DIR / "sample.txt", "rb") as f:
            u = SimpleUploadedFile("sample.txt", f.read())
            resp = client.post(f"/api/threads/{thread.id}/files/", {"files": u}, format="multipart")
            assert resp.status_code == 201

        # Send initial message about the file
        ans, _ = await run_chat(thread, "What file did I upload?")
        print("Initial response:", ans[:60])
        
        # Send 15 intervening short messages
        print("Sending 15 intervening messages...")
        for i in range(1, 16):
            await Message.objects.acreate(thread=thread, role="user", content=f"Count checkpoint {i}")
            await Message.objects.acreate(thread=thread, role="assistant", content=f"Acknowledged {i}.")
            
        # Now ask about the uploaded file
        ans, _ = await run_chat(thread, "What was the exact greeting text in the text file I uploaded at the start of our conversation?")
        print(f"15 messages later response: {ans}")
        assert "hello world" in ans.lower()
        print("[15 messages later] PASS")

async def test_file_generation_all_types(user, client, mock_jwks):
    """Test 4: Generate each output type from a prompt, verify file opens and download works."""
    print("\n--- TEST 4: Generate all output types and test download ---")
    kinds = [
        ("md", "Generate a markdown file named notes.md with title 'Meeting Notes' and content 'Action items discussed.' using generate_file."),
        ("csv", "Generate a CSV file named table.csv with headers ['Product', 'Price'] and rows [['Apple', 2], ['Banana', 1]] using generate_file."),
        ("docx", "Generate a Word document named report.docx with title 'Annual Report' and paragraphs ['The year was successful.'] using generate_file."),
        ("pdf", "Generate a PDF document named summary.pdf with title 'Executive Summary' and paragraphs ['High level overview.'] using generate_file."),
        ("xlsx", "Generate an Excel spreadsheet named finance.xlsx with headers ['Quarter', 'Revenue'] and rows [['Q1', 1000], ['Q2', 1500]] using generate_file."),
        ("pptx", "Generate a PowerPoint presentation named deck.pptx with title 'Pitch Deck' and slides [{'title': 'Problem', 'content': 'Inefficiency in market'}] using generate_file."),
        ("png", "Generate a PNG chart named sales.png with chart_type 'bar', labels ['Jan', 'Feb'], and values [10, 25] using generate_file."),
        ("code", "Generate a python code file named solution.py with code 'def solve(): return 42' using generate_file."),
    ]
    
    with patch("chat.authentication.get_jwks_client", return_value=mock_jwks):
        for kind, prompt in kinds:
            thread = Thread.objects.create(user=user, title=f"Gen {kind} test")
            ans, files = await run_chat(thread, prompt)
            print(f"[{kind}] Agent reply: {ans[:80]}... Files: {files}")
            
            gf = GeneratedFile.objects.filter(thread=thread, kind=kind).order_by("-id").first()
            assert gf is not None, f"GeneratedFile record missing for {kind}"
            assert os.path.exists(gf.path), f"File missing on disk for {kind}: {gf.path}"
            assert os.path.getsize(gf.path) > 0, f"File empty for {kind}"
            
            # Verify file opens validly
            if kind == "md":
                text = open(gf.path, "r", encoding="utf-8").read()
                assert len(text) > 0
            elif kind == "csv":
                import pandas as pd
                df = pd.read_csv(gf.path)
                assert len(df) >= 1
            elif kind == "docx":
                import docx
                d = docx.Document(gf.path)
                assert len(d.paragraphs) >= 1 or len(d.tables) >= 1
            elif kind == "pdf":
                from pypdf import PdfReader
                r = PdfReader(gf.path)
                assert len(r.pages) >= 1
            elif kind == "xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(gf.path)
                assert len(wb.sheetnames) >= 1
            elif kind == "pptx":
                from pptx import Presentation
                p = Presentation(gf.path)
                assert len(p.slides) >= 1
            elif kind == "png":
                from PIL import Image
                im = Image.open(gf.path)
                assert im.size[0] > 0
            elif kind == "code":
                code = open(gf.path, "r", encoding="utf-8").read()
                compile(code, "test.py", "exec")
                
            # Test download
            resp = client.get(f"/api/generated/{gf.id}/download/")
            assert resp.status_code == 200, f"Download failed for {kind}: {resp.status_code}"
            assert resp.get("X-Content-Type-Options") == "nosniff"
            print(f"[{kind}] PASS: Generated valid file and download succeeded.")

async def test_revise_generated_xlsx(user, client, mock_jwks):
    """Test 5: Revise a generated xlsx ("add a column") and verify version 2."""
    print("\n--- TEST 5: Revise generated xlsx ---")
    thread = Thread.objects.create(user=user, title="Revise XLSX test")
    with patch("chat.authentication.get_jwks_client", return_value=mock_jwks):
        # Generate initial version
        await run_chat(
            thread,
            "Generate an Excel file named metrics.xlsx with headers ['Month', 'Users'] and rows [['Jan', 100], ['Feb', 200]] using generate_file."
        )
        v1 = GeneratedFile.objects.filter(thread=thread, name="metrics.xlsx").first()
        assert v1 is not None and v1.version == 1
        
        # Revise to add column
        await run_chat(
            thread,
            f"Please revise metrics.xlsx (based_on {v1.id}) to add a new column 'Churn' with values [5, 8] using generate_file."
        )
        v2 = GeneratedFile.objects.filter(thread=thread, name="metrics.xlsx", version=2).first()
        assert v2 is not None, "Version 2 of metrics.xlsx not generated"
        assert v2.parent_id == v1.id, "Version 2 parent not linked to version 1"
        
        import openpyxl
        wb = openpyxl.load_workbook(v2.path)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        print("Version 2 headers:", headers)
        assert any("churn" in str(h).lower() for h in headers)
        print("[Revise XLSX] PASS: Version 2 successfully created with new column.")

async def test_corrupt_unsupported_and_oversized_files(user, client, mock_jwks):
    """Test 6: Corrupt file, unsupported type, and oversized file give clear errors."""
    print("\n--- TEST 6: Corrupt, unsupported, oversized files ---")
    thread = Thread.objects.create(user=user, title="Errors Test")
    with patch("chat.authentication.get_jwks_client", return_value=mock_jwks):
        # 1. Unsupported type (.exe)
        bad_file = SimpleUploadedFile("malware.exe", b"MZ\x90\x00\x03\x00")
        resp = client.post(f"/api/threads/{thread.id}/files/", {"files": bad_file}, format="multipart")
        assert resp.status_code == 201 # View creates record with status failed
        uf = UploadedFile.objects.filter(thread=thread, name="malware.exe").first()
        assert uf.status == "failed"
        assert "not permitted" in uf.error.lower()
        print("[Unsupported file] PASS: Clear error recorded:", uf.error)

        # 2. Corrupted file (corrupt zip)
        corrupt_zip = SimpleUploadedFile("broken.zip", b"PK\x03\x04randomcorruptedbytes")
        resp = client.post(f"/api/threads/{thread.id}/files/", {"files": corrupt_zip}, format="multipart")
        uf_zip = UploadedFile.objects.filter(thread=thread, name="broken.zip").first()
        assert uf_zip.status == "failed"
        assert len(uf_zip.error) > 0
        print("[Corrupt file] PASS: Clear parsing failure recorded:", uf_zip.error)

        # 3. Oversized file (> 25MB)
        big_content = b"0" * (26 * 1024 * 1024)
        big_file = SimpleUploadedFile("huge.txt", big_content)
        resp = client.post(f"/api/threads/{thread.id}/files/", {"files": big_file}, format="multipart")
        assert resp.status_code in (400, 201)
        if resp.status_code == 400:
            assert "exceeds limit" in str(resp.data).lower()
            print("[Oversized file] PASS: Request rejected with 400 exceeds limit.")
        else:
            uf_big = UploadedFile.objects.filter(thread=thread, name="huge.txt").first()
            assert uf_big.status == "failed"
            assert "exceeds limit" in uf_big.error.lower()
            print("[Oversized file] PASS: UploadedFile status failed with exceeds limit.")

async def test_tool_failures_explained_to_user(user, client, mock_jwks):
    """Test 7: Confirm tool failures are shown to the user, not dropped."""
    print("\n--- TEST 7: Tool failures shown to user ---")
    thread = Thread.objects.create(user=user, title="Tool Failure Test")
    with patch("chat.authentication.get_jwks_client", return_value=mock_jwks):
        ans, _ = await run_chat(thread, "Please run failing_tool with query='trigger_crash'.")
        print(f"Tool failure explanation response: {ans}")
        # The agent should explain the failure or error
        assert any(w in ans.lower() for w in ["fail", "error", "unable", "simulated", "problem", "cannot"])
        print("[Tool failure explained] PASS: Agent explained failure to user.")

async def main():
    print("========================================")
    print("STARTING STEP 4 END-TO-END AUDIT SUITE")
    print("========================================")
    user, _ = User.objects.get_or_create(username="e2e_master_user", defaults={"email": "master@example.com"})
    client, mock_jwks = make_auth_client(user)
    
    await test_file_reading_all_types(user, client, mock_jwks)
    await test_two_files_in_one_message(user, client, mock_jwks)
    await test_ask_15_messages_later(user, client, mock_jwks)
    await test_file_generation_all_types(user, client, mock_jwks)
    await test_revise_generated_xlsx(user, client, mock_jwks)
    await test_corrupt_unsupported_and_oversized_files(user, client, mock_jwks)
    await test_tool_failures_explained_to_user(user, client, mock_jwks)
    
    print("\n========================================")
    print("ALL STEP 4 END-TO-END TESTS PASSED!")
    print("========================================")

if __name__ == "__main__":
    asyncio.run(main())
