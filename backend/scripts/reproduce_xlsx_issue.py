import os
import sys
import json
import time
import asyncio
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
os.environ["DJANGO_ALLOW_ASYNC_UNSAFE"] = "true"
import django
django.setup()

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from chat.models import Thread, Message, UploadedFile, GeneratedFile
from chat.streaming import stream_chat_turn
from filesvc.service import file_service
from filesvc.storage import storage

User = get_user_model()
SAMPLE_XLSX = BASE_DIR / "sample_files" / "sales.xlsx"

def inspect_workbook(file_path: str, label: str):
    print(f"\n{'='*25} INSPECTING {label} {'='*25}")
    print(f"File path: {file_path}")
    if not os.path.exists(file_path):
        print(f"ERROR: File {file_path} does not exist!")
        return

    # 1. Default load (formula strings)
    wb_default = openpyxl.load_workbook(file_path, data_only=False)
    sheet_default = wb_default.active

    # 2. data_only=True load (cached values)
    wb_data_only = openpyxl.load_workbook(file_path, data_only=True)
    sheet_data_only = wb_data_only.active

    max_r = sheet_default.max_row
    max_c = sheet_default.max_column
    print(f"Dimensions: {max_r} rows x {max_c} cols")

    print("\n--- Last 5 Rows (Default vs data_only=True) ---")
    start_r = max(1, max_r - 4)
    for r in range(start_r, max_r + 1):
        def_vals = [sheet_default.cell(row=r, column=c).value for c in range(1, max_c + 1)]
        cached_vals = [sheet_data_only.cell(row=r, column=c).value for c in range(1, max_c + 1)]
        print(f"Row {r:2d} | DEFAULT  : {def_vals}")
        print(f"Row {r:2d} | DATA_ONLY: {cached_vals}")

    print("\n--- Formula Cells Analysis ---")
    formula_count = 0
    empty_cached_count = 0
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            val_def = sheet_default.cell(row=r, column=c).value
            val_data = sheet_data_only.cell(row=r, column=c).value
            is_formula = isinstance(val_def, str) and val_def.startswith("=")
            if is_formula:
                formula_count += 1
                coord = sheet_default.cell(row=r, column=c).coordinate
                print(f"Cell {coord:4s}: Formula = '{val_def}' | Cached (data_only) = {repr(val_data)}")
                if val_data is None:
                    empty_cached_count += 1

    print(f"Total formula cells: {formula_count}, Formula cells with None cached value: {empty_cached_count}")
    print("=" * 60)

async def run_prompt(thread, user, prompt_text: str):
    print(f"\n[User Query]: '{prompt_text}'")
    events = []
    generated_file_ids = []
    assistant_text = ""
    await Message.objects.acreate(thread=thread, role="user", content=prompt_text, status="complete")
    async for raw in stream_chat_turn(thread=thread, user_prompt=prompt_text):
        events.append(raw)
        lines = raw.strip().split("\n")
        ev_name = None
        ev_data = None
        for line in lines:
            if line.startswith("event: "):
                ev_name = line[7:].strip()
            elif line.startswith("data: "):
                ev_data = line[6:].strip()

        if ev_name == "token" and ev_data:
            try:
                parsed_token = json.loads(ev_data)
                t = parsed_token.get("content", "") or parsed_token.get("token", "")
                assistant_text += t
            except:
                pass
        elif ev_name == "file" and ev_data:
            try:
                f_info = json.loads(ev_data)
                generated_file_ids.append(f_info.get("id"))
                print(f"  [File Event]: {f_info.get('name')} (id={f_info.get('id')}, v={f_info.get('version')})")
            except:
                pass
        elif ev_name == "status" and ev_data:
            print(f"  [Status Event]: {ev_data}")

    print(f"\n[Assistant Response]:\n{assistant_text.strip()}\n")
    return generated_file_ids

async def main():
    print("=" * 70)
    print("REPRODUCING XLSX GENERATION ISSUE")
    print("=" * 70)

    # 1. Setup test user and thread
    user, _ = await User.objects.aget_or_create(username="xlsx_tester", defaults={"email": "xlsx@example.com"})
    thread = await Thread.objects.acreate(user=user, title="XLSX Revenue Test")

    # 2. Upload sales.xlsx
    with open(SAMPLE_XLSX, "rb") as f:
        file_obj = SimpleUploadedFile("sales.xlsx", f.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        saved_path, written_size = storage.save_raw(file_obj, ".xlsx")
        uf = await UploadedFile.objects.acreate(
            thread=thread,
            name="sales.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size=written_size,
            path=saved_path,
            status="pending",
        )
        uf = await asyncio.to_thread(file_service.parse_file_sync, uf.id)

    print(f"Uploaded and parsed sales.xlsx (UploadedFile ID: {uf.id})")
    print(f"File summary: {uf.summary[:150]}...")

    # 3. Prompt 1: Create an Excel file with monthly revenue from spreadsheet and add a total row
    p1 = "Create an Excel file with the monthly revenue from the spreadsheet and add a total row."
    fids_1 = await run_prompt(thread, user, p1)

    gf_1 = await GeneratedFile.objects.filter(thread=thread).order_by("-id").afirst()
    if gf_1:
        inspect_workbook(gf_1.path, f"Prompt 1 Output ({gf_1.name} v{gf_1.version})")
    else:
        print("ERROR: No generated file created for prompt 1!")

    # 4. Prompt 2: Add a column showing the month-over-month growth percentage
    p2 = "Add a column showing the month-over-month growth percentage."
    fids_2 = await run_prompt(thread, user, p2)

    gf_2 = await GeneratedFile.objects.filter(thread=thread).order_by("-id").afirst()
    if gf_2:
        inspect_workbook(gf_2.path, f"Prompt 2 Output ({gf_2.name} v{gf_2.version})")
        from django.test import AsyncClient
        c = AsyncClient()
        token = f"service-token-{user.username}"
        prev_res = await c.get(f"/api/generated/{gf_2.id}/preview/", headers={"Authorization": f"Bearer {token}"})
        print("\n--- PREVIEW ENDPOINT RESULT ---")
        print(f"Status: {prev_res.status_code}")
        import json as j_mod
        print(f"Preview Payload:\n{j_mod.dumps(prev_res.json(), indent=2)}")
    else:
        print("ERROR: No generated file created for prompt 2!")

if __name__ == "__main__":
    asyncio.run(main())
